import csv
import json
import logging
from pathlib import Path

from invoice_sorter.models import Document

logger = logging.getLogger(__name__)


class Exporter:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
    
    def export_json(self, documents: list[Document], filename: str = "invoices.json") -> Path:
        filepath = self.output_dir / filename
        data = [doc.to_dict() for doc in documents]
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        logger.info(f"Exported {len(documents)} documents to {filepath}")
        return filepath
    
    def export_csv(self, documents: list[Document], filename: str = "invoices.csv") -> Path:
        filepath = self.output_dir / filename
        if not documents:
            logger.warning("No documents to export")
            return filepath
        
        fieldnames = [
            "file_path", "document_type", "vendor", "amount", "currency",
            "date", "invoice_number", "category", "description", "confidence",
            "needs_review"
        ]
        
        with open(filepath, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for doc in documents:
                writer.writerow(doc.to_dict())
        
        logger.info(f"Exported {len(documents)} documents to {filepath}")
        return filepath
    
    def export_excel(self, documents: list[Document], filename: str = "invoices.xlsx") -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.error("openpyxl not installed, cannot export Excel")
            raise
        
        filepath = self.output_dir / filename
        wb = Workbook()
        ws = wb.active
        ws.title = "Invoices & Receipts"
        
        headers = [
            "File Path", "Type", "Vendor", "Amount", "Currency",
            "Date", "Invoice Number", "Category", "Description", "Confidence"
        ]
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        
        for row, doc in enumerate(documents, 2):
            ws.cell(row=row, column=1, value=doc.file_path)
            ws.cell(row=row, column=2, value=doc.document_type.value)
            ws.cell(row=row, column=3, value=doc.vendor)
            ws.cell(row=row, column=4, value=doc.amount)
            ws.cell(row=row, column=5, value=doc.currency)
            ws.cell(row=row, column=6, value=doc.date.isoformat() if doc.date else "")
            ws.cell(row=row, column=7, value=doc.invoice_number)
            ws.cell(row=row, column=8, value=doc.category.value)
            ws.cell(row=row, column=9, value=doc.description)
            ws.cell(row=row, column=10, value=doc.confidence)
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
        
        wb.save(filepath)
        logger.info(f"Exported {len(documents)} documents to {filepath}")
        return filepath
    
    def export_summary(self, summary: dict, filename: str = "summary.json") -> Path:
        filepath = self.output_dir / filename
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(summary, f, indent=2, ensure_ascii=False)
        logger.info(f"Exported summary to {filepath}")
        return filepath


def export_all(
    documents: list[Document],
    output_dir: Path,
    formats: list[str] | None = None,
) -> dict[str, Path]:
    if formats is None:
        formats = ["json", "csv"]
    
    exporter = Exporter(output_dir)
    results = {}
    
    if "json" in formats:
        results["json"] = exporter.export_json(documents)
    if "csv" in formats:
        results["csv"] = exporter.export_csv(documents)
    if "excel" in formats:
        results["excel"] = exporter.export_excel(documents)
    
    return results